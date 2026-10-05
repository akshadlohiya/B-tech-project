#!/usr/bin/env python3
"""Generate the UWB RTLS project presentation as a .pptx in the COEP Tech theme.

Reproduces the uploaded Beamer format: blue header bar with the frame title,
a section-navigation strip, left-pointing triangle bullets, and a footer with
the COEP logo (left), institute name (centre) and slide number (right).
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "COEP_Tech_University_Logo.jpg")
SHOT = os.path.join(HERE, "twin_screenshot.png")
ARCH = os.path.join(HERE, "architecture.png")
CHART = os.path.join(HERE, "accuracy_chart.png")
UML = os.path.join(HERE, "uml_class.png")

# --- palette (COEP beamer blue) ------------------------------------------------
BLUE      = RGBColor(0x2A, 0x2A, 0xA8)   # title-bar / structure blue
BLUE_DK   = RGBColor(0x1F, 0x1F, 0x7A)
NAVY_TXT  = RGBColor(0x22, 0x22, 0x88)
GREY      = RGBColor(0x55, 0x55, 0x55)
BLACK     = RGBColor(0x22, 0x22, 0x22)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
RED       = RGBColor(0xB0, 0x22, 0x22)

SECTIONS = ["Introduction", "Literature Review", "Research Gap", "Problem Statement",
            "Key Concepts", "Work Done", "Work Plan", "Publication & References"]

prs = Presentation()
prs.slide_width  = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

TRI = "◀"   # ◀

_slide_no = {"n": 0}
TOTAL = {"n": 0}


def _txbox(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(1)
    return tb, tf


def _set(run, text, size=18, bold=False, italic=False, color=BLACK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font


def _rect(slide, l, t, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
    sp.shadow.inherit = False
    return sp


def _round(slide, l, t, w, h, fill):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    sp.shadow.inherit = True
    return sp


def chrome(slide, title, active_section):
    """Header nav strip + blue title bar + footer."""
    # --- top navigation strip ---
    _rect(slide, 0, 0, SW, Inches(0.34), RGBColor(0xE9, 0xE9, 0xF4))
    n = len(SECTIONS)
    seg = SW / n
    for i, name in enumerate(SECTIONS):
        tb, tf = _txbox(slide, Emu(int(i * seg)), Inches(0.03), Emu(int(seg)), Inches(0.28))
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        _set(r, name, size=9, bold=(name == active_section),
             color=(BLUE_DK if name == active_section else GREY))
    # --- blue title bar ---
    _rect(slide, 0, Inches(0.34), SW, Inches(0.92), BLUE)
    tb, tf = _txbox(slide, Inches(0.35), Inches(0.42), SW - Inches(0.7), Inches(0.78))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run(); _set(r, title, size=26, bold=True, color=WHITE)
    # --- footer ---
    if os.path.exists(LOGO):
        slide.shapes.add_picture(LOGO, Inches(0.12), SH - Inches(0.62), height=Inches(0.5))
    tb, tf = _txbox(slide, Inches(2.4), SH - Inches(0.6), SW - Inches(4.0), Inches(0.5))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); _set(r, "COEP Technological University [COEP Tech]", size=12, color=NAVY_TXT)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); _set(r2, "(A Unitary Public University of Government of Maharashtra)", size=8, color=NAVY_TXT)
    _slide_no["n"] += 1
    tb, tf = _txbox(slide, SW - Inches(1.4), SH - Inches(0.5), Inches(1.25), Inches(0.35))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); _set(r, f"{_slide_no['n']} / {TOTAL['n']}", size=11, bold=True, color=GREY)


def body(slide, top=Inches(1.45), left=Inches(0.6), width=None, height=None):
    if width is None:
        width = SW - Inches(1.2)
    if height is None:
        height = SH - top - Inches(0.75)
    tb, tf = _txbox(slide, left, top, width, height)
    return tf


def bullets(tf, items, size=18, first=True):
    """items: list of (text, level, kwargs)."""
    for idx, item in enumerate(items):
        if isinstance(item, tuple):
            text, level = item[0], item[1]
            kw = item[2] if len(item) > 2 else {}
        else:
            text, level, kw = item, 0, {}
        p = tf.paragraphs[0] if (idx == 0 and first) else tf.add_paragraph()
        p.level = level
        p.space_after = Pt(6)
        marker = kw.get("marker", TRI)
        if marker:
            rb = p.add_run()
            _set(rb, marker + "  ", size=size - 4, color=kw.get("mcolor", BLUE))
        # allow inline segments: kw['segs'] = [(text,bold,color,italic),...]
        if "segs" in kw:
            for (t, b, c, i) in kw["segs"]:
                r = p.add_run(); _set(r, t, size=size, bold=b, italic=i, color=c)
        else:
            r = p.add_run()
            _set(r, text, size=size, bold=kw.get("bold", False),
                 italic=kw.get("italic", False), color=kw.get("color", BLACK))


def block(slide, l, t, w, h, heading, lines, hsize=15, bsize=14, fill=RGBColor(0xEC,0xEC,0xF7)):
    """A titled rounded box (like a beamer block)."""
    _round(slide, l, t, w, Inches(0.42), BLUE)
    tb, tf = _txbox(slide, l + Inches(0.12), t + Inches(0.02), w - Inches(0.24), Inches(0.4))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run(); _set(r, heading, size=hsize, bold=True, color=WHITE)
    bx = _rect(slide, l, t + Inches(0.42), w, h - Inches(0.42), fill)
    tb, tf = _txbox(slide, l + Inches(0.15), t + Inches(0.5), w - Inches(0.3), h - Inches(0.55))
    if isinstance(lines, str):
        r = tf.paragraphs[0].add_run(); _set(r, lines, size=bsize, color=BLACK)
    else:
        bullets(tf, lines, size=bsize)
    return bx


def new(title, section):
    s = prs.slides.add_slide(BLANK)
    chrome(s, title, section)
    return s

# We need TOTAL slide count for "x / y"; count by pre-building a list of builders.
builders = []

# ------------------------------------------------------------------ TITLE
def s_title():
    s = prs.slides.add_slide(BLANK)
    _slide_no["n"] += 1
    # rounded blue title block
    _round(s, Inches(0.9), Inches(1.4), SW - Inches(1.8), Inches(1.9), BLUE)
    tb, tf = _txbox(s, Inches(1.2), Inches(1.55), SW - Inches(2.4), Inches(1.6))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); _set(r, "UWB Indoor Real-Time Location System", size=30, bold=True, color=WHITE)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r = p2.add_run(); _set(r, "with a Live 3D Digital Twin", size=30, bold=True, color=WHITE)
    p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER
    r = p3.add_run(); _set(r, "B. Tech. Project  —  Computer Science & Engineering", size=15, italic=True, color=RGBColor(0xDD,0xDD,0xF5))
    # authors
    tb, tf = _txbox(s, Inches(1.2), Inches(3.7), SW - Inches(2.4), Inches(1.2))
    for i, (nm, role) in enumerate([
        ("Nahush Kale      ·      Akshad Lohiya      ·      Purnank Vasaikar", ""),
        ("Under the guidance of  Prof. Dr. H. D. Gadade", "g"),
        ("Department of Computer Science & Engineering,  COEP Technological University, Pune", "d"),
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        _set(r, nm, size=(17 if i == 0 else 14), bold=(i == 0),
             italic=(role == "g"), color=(BLACK if i == 0 else GREY))
    if os.path.exists(LOGO):
        s.shapes.add_picture(LOGO, int(SW/2 - Inches(0.7)), Inches(5.15), height=Inches(1.15))
    tb, tf = _txbox(s, Inches(1.2), Inches(6.35), SW - Inches(2.4), Inches(0.5))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); _set(r, "September 22, 2026", size=14, bold=True, color=NAVY_TXT)
builders.append(s_title)

# ------------------------------------------------------------------ OUTLINE
def s_outline():
    s = new("Outline", None)
    tf = body(s)
    items = [(f"{i+1}.  {name}", 0, {"marker": "", "bold": True, "color": BLUE_DK})
             for i, name in enumerate(SECTIONS)]
    bullets(tf, items, size=20)
builders.append(s_outline)

# ------------------------------------------------------------------ INTRODUCTION
def s_intro():
    s = new("Introduction", "Introduction")
    tf = body(s)
    bullets(tf, [
        ("", 0, {"segs":[("UWB (Ultra-Wideband): ",True,BLACK,False),
                 ("very short, very wide-band radio pulses ⇒ fine time resolution + strong multipath rejection ⇒ centimetre-level ranging indoors.",False,BLACK,False)]}),
        ("", 0, {"segs":[("RTLS: ",True,BLACK,False),
                 ("Real-Time Location System — track tagged assets & people live, indoors, where GPS cannot reach.",False,BLACK,False)]}),
        ("", 0, {"segs":[("Digital Twin: ",True,BLACK,False),
                 ("a live, labelled 3D model of the floor that shows every tracked item moving in real time.",False,BLACK,False)]}),
        ("", 0, {"segs":[("Our system: ",True,BLACK,False),
                 ("ceiling UWB anchors range the tags via DS-TWR → a position engine turns ranges into a smooth track → a browser twin renders room, anchors and tags.",False,BLACK,False)]}),
    ], size=19)
builders.append(s_intro)

def s_whyuwb():
    s = new("Why UWB Indoors", "Introduction")
    tf = body(s)
    bullets(tf, [
        ("", 0, {"segs":[("GPS is blocked indoors; ",False,BLACK,False),("Wi-Fi / BLE give metre-to-tens-of-metre error",True,BLACK,False),(" (narrowband, multipath-prone).",False,BLACK,False)]}),
        "UWB is the leading technology for high-accuracy indoor positioning.",
        ("", 0, {"segs":[("Target application: ",True,BLACK,False),("warehouse-scale asset & personnel tracking — inventory location, worker safety, geofencing / zone alerts.",False,BLACK,False)]}),
        ("", 0, {"segs":[("Hardware: ",True,BLACK,False),("Makerfabs MaUWB_DW3000 nodes; firmware performs DS-TWR ranging and reports a distance per anchor plus signal-power diagnostics.",False,BLACK,False)]}),
    ], size=19)
builders.append(s_whyuwb)

# ------------------------------------------------------------------ LITERATURE
def s_lit():
    s = new("Literature Review", "Literature Review")
    half = (SW - Inches(1.5)) / 2
    # left
    tf = body(s, left=Inches(0.6), width=half)
    r = tf.paragraphs[0].add_run(); _set(r, "Key studies", size=17, bold=True, color=BLUE_DK)
    bullets(tf, [
        ("Paszek et al., Sensors 2021 — UWB LOS/NLOS simulator + accuracy analysis.", 0),
        ("Appl. Sci. 2025 — NLOS-exclusion positioning (RMSE 0.124 m, ~24% ↑).", 0),
        ("Appl. Sci. 2024 — UWB RTLS deployment & error study.", 0),
        ("Zhao et al., IJRR 2024 — UTIL dataset: raw TDOA + power-diff + mm Vicon truth.", 0),
        ("Alexandria Eng. J. 2018 — UWB error modelling.", 0),
    ], size=15, first=False)
    # right
    tf = body(s, left=Inches(0.75) + half, width=half)
    r = tf.paragraphs[0].add_run(); _set(r, "Takeaways", size=17, bold=True, color=BLUE_DK)
    bullets(tf, [
        ("NLOS-aware weighting is the single biggest accuracy lever.", 0),
        ("A hardware-faithful simulator lets the full pipeline be built before hardware arrives.", 0),
        ("Real mm-ground-truth data (UTIL) enables noise/NLOS calibration and EKF validation.", 0),
        ("DS-TWR cancels clock drift ⇒ no anchor time-synchronisation needed.", 0),
    ], size=15, first=False)
builders.append(s_lit)

def s_util():
    s = new("Reference Dataset: UTIL", "Literature Review")
    tf = body(s)
    bullets(tf, [
        ("UofT UTIAS DSL, IJRR 2024 — Decawave DWM1000; 4 anchor constellations, ~150 min of flights.", 0),
        ("Provides raw UWB TDOA, SNR / power-difference (LOS/NLOS labels), IMU, altitude, mm-accurate Vicon ground truth, plus a reference EKF and parsers.", 0),
        ("", 0, {"segs":[("Cross-radio caveat: ",True,RED,False),("DWM1000 ≠ our DW3000 ⇒ a fitted model is a prior needing DW3000 recalibration, not a direct transplant.",False,BLACK,False)]}),
    ], size=19)
builders.append(s_util)

# ------------------------------------------------------------------ RESEARCH GAP
def s_gap():
    s = new("Research Gap", "Research Gap")
    tf = body(s)
    bullets(tf, [
        ("", 0, {"segs":[("Many demos shortcut: ",True,BLACK,False),("they display measured ranges but stream the true position to the UI ⇒ no real solve, no NLOS handling, fake accuracy.",False,BLACK,False)]}),
        "NLOS down-weighting is usually hand-tuned constants, not fitted to real error data.",
        "Simulator and hardware are typically separate codebases ⇒ costly, error-prone integration.",
        "Weak height (z) observability with co-planar ceiling anchors is rarely addressed.",
    ], size=18)
    y = Inches(4.7)
    block(s, Inches(0.6), y, SW - Inches(1.2), Inches(1.9),
          "Our gap-closing contributions",
          [("A pipeline that actually estimates position from noisy, partly-NLOS ranges.", 0),
           ("A data-driven NLOS / noise model fitted on real mm-truth UWB data.", 0),
           ("One codebase — simulator ↔ hardware behind a single interface.", 0)],
          bsize=15)
builders.append(s_gap)

# ------------------------------------------------------------------ PROBLEM
def s_problem():
    s = new("Problem Statement", "Problem Statement")
    block(s, Inches(0.6), Inches(1.5), SW - Inches(1.2), Inches(1.15), "Goal",
          "Build an end-to-end indoor RTLS that ranges tags from ceiling anchors, solves each tag's "
          "position from noisy / partly-NLOS ranges, smooths the estimate over time, and visualises it "
          "live on a 3D digital twin.", bsize=15)
    tf = body(s, top=Inches(2.95))
    bullets(tf, [
        ("", 0, {"segs":[("The estimate must be honest: ",True,BLACK,False),("position is computed from measurements, not read from ground truth — and accuracy is reported against truth.",False,BLACK,False)]}),
        "Architected so the same software runs against a realistic simulator today and real UWB hardware tomorrow.",
        ("", 0, {"segs":[("Hardware integration confined to a single drop-in source file ",False,BLACK,False),("(SimAnchorSource → HardwareAnchorSource).",False,GREY,True)]}),
    ], size=18)
builders.append(s_problem)

def s_obj():
    s = new("Objectives", "Problem Statement")
    tf = body(s)
    items = [
        "Define a single versioned data contract shared by simulator & hardware.",
        "Build a hardware-faithful DS-TWR simulator (LOS/NLOS, antenna-delay bias, power diagnostics, drops).",
        "Implement a multilateration solver with NLOS down-weighting from the DW3000 power gap.",
        "Fuse fixes with a constant-velocity Kalman filter ⇒ smooth track + velocity + uncertainty.",
        "Render a config-driven 3D digital twin (trails, uncertainty rings, zones, live accuracy HUD).",
        "Calibrate & validate the noise / NLOS model on real mm-truth UWB data (UTIL).",
        "Keep hardware integration to a single file.",
    ]
    bullets(tf, [(f"{i+1}.  {t}", 0, {"marker":"", "bold":False}) for i, t in enumerate(items)], size=18)
builders.append(s_obj)

# ------------------------------------------------------------------ KEY CONCEPTS
def s_kc1():
    s = new("Key Concepts (1/4) — UWB Positioning Basics", "Key Concepts")
    block(s, Inches(0.6), Inches(1.5), SW - Inches(1.2), Inches(1.05), "UWB — Ultra-Wideband",
          "Radio using very short pulses over a very wide bandwidth (>500 MHz) ⇒ fine time resolution "
          "+ strong multipath rejection ⇒ centimetre-level ranging.", bsize=15)
    half = (SW - Inches(1.5)) / 2
    tf = body(s, top=Inches(2.85), left=Inches(0.6), width=half)
    bullets(tf, [
        ("", 0, {"segs":[("RTLS: ",True,BLACK,False),("live indoor tracking of tagged items.",False,BLACK,False)]}),
        ("", 0, {"segs":[("Anchor: ",True,BLACK,False),("fixed node of known coordinates.",False,BLACK,False)]}),
        ("", 0, {"segs":[("Tag: ",True,BLACK,False),("mobile node on the tracked asset / person.",False,BLACK,False)]}),
    ], size=16, first=False)
    tf = body(s, top=Inches(2.85), left=Inches(0.75)+half, width=half)
    bullets(tf, [
        ("", 0, {"segs":[("LOS: ",True,BLACK,False),("Line-of-Sight — clear direct path; small error.",False,BLACK,False)]}),
        ("", 0, {"segs":[("NLOS: ",True,BLACK,False),("path blocked; signal reflects, so the range reads longer — a positive bias, not just noise.",False,BLACK,False)]}),
    ], size=16, first=False)
builders.append(s_kc1)

def s_kc2():
    s = new("Key Concepts (2/4) — Ranging & Observation Models", "Key Concepts")
    block(s, Inches(0.6), Inches(1.5), SW - Inches(1.2), Inches(1.15),
          "TWR / DS-TWR — (Double-Sided) Two-Way Ranging",
          "Anchor and tag exchange timed messages to measure a distance. Double-sided adds a message so "
          "device clock-drift cancels ⇒ no anchor time-sync needed. This is exactly what our DW3000 hardware reports.",
          bsize=14)
    half = (SW - Inches(1.5)) / 2
    tf = body(s, top=Inches(2.95), left=Inches(0.6), width=half)
    bullets(tf, [
        ("", 0, {"segs":[("TOA: ",True,BLACK,False),("Time-of-Arrival — absolute travel time → a range.",False,BLACK,False)]}),
        ("", 0, {"segs":[("TDOA: ",True,BLACK,False),("Time-Difference-of-Arrival — synchronised anchors compare arrival times; each pair → a hyperbola.",False,BLACK,False)]}),
    ], size=16, first=False)
    block(s, Inches(0.75)+half, Inches(2.95), half, Inches(1.7), "Range vs TDOA",
          "Range needs no sync but a direct distance; TDOA scales to many tags but needs tightly synced "
          "anchors. Our pipeline is range-based; a walled-off module handles TDOA.", bsize=13)
builders.append(s_kc2)

def s_kc3():
    s = new("Key Concepts (3/4) — Position Solving", "Key Concepts")
    block(s, Inches(0.6), Inches(1.5), SW - Inches(1.2), Inches(1.05),
          "Multilateration / Trilateration",
          "Recover a tag's position from ≥ 3 anchor ranges by intersecting range circles / spheres "
          "— solved as a least-squares fit when measurements are noisy.", bsize=15)
    tf = body(s, top=Inches(2.85))
    bullets(tf, [
        ("", 0, {"segs":[("WLS — Weighted Least Squares: ",True,BLACK,False),("each anchor's residual is weighted by its trustworthiness, so unreliable links influence the solution less.",False,BLACK,False)]}),
        ("", 0, {"segs":[("Power diagnostics: ",True,BLACK,False),("DW3000 reports received power (rxPower) and first-path power (fpPower); a large rxPower−fpPower gap ⇒ attenuated first path ⇒ likely NLOS.",False,BLACK,False)]}),
        ("", 0, {"segs":[("NLOS down-weighting: ",True,BLACK,False),("that power gap sets the WLS weight, so a biased (long) NLOS range cannot drag the fix.",False,BLACK,False)]}),
    ], size=17, first=False)
builders.append(s_kc3)

def s_kc4():
    s = new("Key Concepts (4/4) — State Estimation & Visualisation", "Key Concepts")
    half = (SW - Inches(1.5)) / 2
    block(s, Inches(0.6), Inches(1.5), half, Inches(1.25), "Kalman Filter (KF)",
          "A recursive predict → update estimator that fuses a motion model with noisy measurements to "
          "output a smoothed state and its covariance (uncertainty).", bsize=13)
    tf = body(s, top=Inches(2.95), left=Inches(0.6), width=half)
    bullets(tf, [
        ("", 0, {"segs":[("CV model: ",True,BLACK,False),("constant-velocity motion, state [x, z, vx, vz].",False,BLACK,False)]}),
        ("", 0, {"segs":[("EKF — Extended KF: ",True,BLACK,False),("a KF for nonlinear measurements (e.g. TDOA), linearised at each step.",False,BLACK,False)]}),
    ], size=15, first=False)
    block(s, Inches(0.75)+half, Inches(1.5), half, Inches(1.25), "Digital Twin",
          "A live, labelled 3D model of the physical floor rendering every tracked tag, its uncertainty "
          "ellipse, trail and zone in real time.", bsize=13)
    tf = body(s, top=Inches(2.95), left=Inches(0.75)+half, width=half)
    bullets(tf, [
        ("", 0, {"segs":[("Covariance / uncertainty ellipse: ",True,BLACK,False),("how sure the filter is of a position.",False,BLACK,False)]}),
        ("", 0, {"segs":[("Ground truth: ",True,BLACK,False),("a reference (mm Vicon truth in UTIL) to measure error against.",False,BLACK,False)]}),
    ], size=15, first=False)
builders.append(s_kc4)

# ------------------------------------------------------------------ WORK DONE
def s_wd_arch():
    s = new("Work Done — System Architecture", "Work Done")
    if os.path.exists(ARCH):
        picw = Inches(9.2)
        s.shapes.add_picture(ARCH, int((SW - picw) / 2), Inches(1.45), width=picw)
    tf = body(s, top=Inches(4.55))
    bullets(tf, [
        ("", 0, {"segs":[("Anchor Source ",True,BLACK,False),("emits per-anchor range + power diagnostics; ",False,BLACK,False),("Location Engine ",True,BLACK,False),("solves + filters; ",False,BLACK,False),("Backend ",True,BLACK,False),("relays + serves /api/site; ",False,BLACK,False),("Frontend ",True,BLACK,False),("visualises the estimated state.",False,BLACK,False)]}),
        ("", 0, {"segs":[("Shared site config ",True,BLACK,False),("(site.json, metres): anchors, walls+material, zones, tag routes, ranging/noise model — every value provenance-tagged.",False,BLACK,False)]}),
    ], size=15)
builders.append(s_wd_arch)

def s_wd_uml():
    s = new("Work Done — UML Class Diagram", "Work Done")
    if os.path.exists(UML):
        pic = s.shapes.add_picture(UML, Inches(2), Inches(1.42), height=Inches(5.4))
        pic.left = int((SW - pic.width) / 2)
builders.append(s_wd_uml)

def s_wd_data():
    s = new("Work Done — Data Contracts", "Work Done")
    tb, tf = _txbox(s, Inches(0.6), Inches(1.4), SW - Inches(1.2), Inches(0.4))
    r = tf.paragraphs[0].add_run(); _set(r, "One versioned schema is what makes simulator and hardware interchangeable.", size=15, italic=True, color=GREY)
    half = (SW - Inches(1.5)) / 2
    block(s, Inches(0.6), Inches(1.95), half, Inches(2.7),
          "Reading (source → engine) · uwb.reading/1",
          [("range — already contains bias + noise + any NLOS excess", 0),
           ("rxPower, fpPower — the NLOS diagnostic", 0),
           ("nlos, valid, timestamp", 0),
           ("", 0, {"segs":[("No ground truth reaches the solver.",True,RED,False)]})], bsize=14)
    block(s, Inches(0.75)+half, Inches(1.95), half, Inches(2.7),
          "Estimate (engine → twin)",
          [("pos, vel, cov — filtered state + uncertainty", 0),
           ("raw — pre-Kalman fix, for comparison", 0),
           ("truth — sim only, drives the accuracy HUD", 0),
           ("zone, alerts", 0)], bsize=14)
builders.append(s_wd_data)

def s_wd_site():
    s = new("Work Done — Site Config & Mapping Platform", "Work Done")
    half = (SW - Inches(1.5)) / 2
    block(s, Inches(0.6), Inches(1.5), half, Inches(2.1),
          "site.json — single source of truth (metres)",
          [("Anchors, walls (+material), labelled zones, tag routes, ranging/noise + Kalman model; every value provenance-tagged.", 0),
           ("Read by both simulator and twin — one edit updates the whole system.", 0),
           ("Coordinates in local metres, not GPS lat/lon.", 0)], bsize=13)
    block(s, Inches(0.75)+half, Inches(1.5), half, Inches(2.1),
          "Mapping platform (Mappedin-style)",
          [("Map library: create / list / delete maps, each with an unguessable share key.", 0),
           ("2D editor: draw walls / zones, place anchors, lay out routes.", 0),
           ("Save & Apply hot-reloads sim, solver & twin — no restart.", 0)], bsize=13)
    tb, tf = _txbox(s, Inches(0.6), Inches(4.0), SW - Inches(1.2), Inches(0.5))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); _set(r, "Stack:  Python (NumPy / SciPy)  ·  Node.js + Socket.IO  ·  React + Three.js  ·  pytest",
                          size=14, bold=True, color=NAVY_TXT)
builders.append(s_wd_site)

def s_wd_map():
    s = new("Work Done — Map Building (Digital-Twin Platform)", "Work Done")
    half = (SW - Inches(1.5)) / 2
    block(s, Inches(0.6), Inches(1.5), half, Inches(1.85), "Build a map — 2D floor-plan editor",
          [("Draw walls with a material (metal / wood / light — sets NLOS bias) and zones (Dock, Aisle, Staging).", 0),
           ("Place anchors (x, y, z + antenna-delay bias) and lay out tag routes (waypoints, speed, allowed zones).", 0)],
          bsize=12)
    block(s, Inches(0.6), Inches(3.6), half, Inches(1.6), "Storage & sharing",
          [("Maps live in the backend store (backend/maps/), each with an unguessable share key.", 0),
           ("Open ?map=KEY to view / share; gateway runs one map: MAP=<key> python main.py.", 0)],
          bsize=12)
    block(s, Inches(0.75) + half, Inches(1.5), half, Inches(1.4), "Hot-reload flow (no restart)",
          "Edit-Map  →  PUT /api/site  →  backend broadcasts site_updated  →  simulator, solver & 3D twin all reload live.",
          bsize=12)
    block(s, Inches(0.75) + half, Inches(3.15), half, Inches(2.05), "Single source of truth — site.json (uwb.site/1)",
          [("Floor, anchors, walls, zones, tags, ranging / noise model in local metres.", 0),
           ("Read by both simulator and twin; every value provenance-tagged.", 0),
           ("Three shipped maps: demo, util, wh-tdoa.", 0)],
          bsize=12)
builders.append(s_wd_map)

def s_wd_sim():
    s = new("Work Done — Hardware-Faithful Simulator", "Work Done")
    tb, tf = _txbox(s, Inches(0.6), Inches(1.4), SW - Inches(1.2), Inches(0.4))
    r = tf.paragraphs[0].add_run(); _set(r, "SimAnchorSource emits exactly what the DW3000 firmware reports (per anchor, per tag, at 5 Hz):", size=15, italic=True, color=GREY)
    tf = body(s, top=Inches(1.95))
    bullets(tf, [
        ("Antenna-delay bias (fixed per-anchor offset) + Gaussian range noise.", 0),
        ("LOS/NLOS decision via a 2D wall-occlusion test.", 0),
        ("", 0, {"segs":[("Positive-only NLOS bias ",True,BLACK,False),("— a blocked path is always reported longer, magnitude drawn per obstructing material (metal racks bias more than wood).",False,BLACK,False)]}),
        ("Power diagnostics (rxPower, fpPower) reproducing the DW3000 gap.", 0),
        ("Dropped readings (higher probability under NLOS); waypoint motion on a simulated clock.", 0),
    ], size=17)
builders.append(s_wd_sim)

def s_wd_solver():
    s = new("Work Done — Multilateration Solver", "Work Done")
    tf = body(s)
    bullets(tf, [
        ("2D trilateration with height compensation; minimise the weighted range residual", 0),
        ("min_p  Σ w_i ( ‖p − a_i‖ − r_i )²    via scipy least_squares (Levenberg–Marquardt), seeded from centroid / last position.", 1, {"marker":"", "italic":True, "color":BLUE_DK}),
        ("", 0, {"segs":[("NLOS down-weighting: ",True,BLACK,False),("a large rxPower−fpPower gap lowers a link's weight w_i, so a biased (long) NLOS range cannot drag the fix.",False,BLACK,False)]}),
        ("", 0, {"segs":[("Verified: ",True,BLACK,False),("recovers a known point to < 1 cm; beats naively trusting a biased anchor.",False,BLACK,False)]}),
    ], size=18)
builders.append(s_wd_solver)

def s_wd_kf():
    s = new("Work Done — Kalman Filter + Digital Twin", "Work Done")
    lw = Inches(6.6)
    tf = body(s, left=Inches(0.6), width=lw)
    bullets(tf, [
        ("", 0, {"segs":[("CV Kalman per tag, ",True,BLACK,False),("state [x, z, vx, vz]: predict(dt) → update with the raw fix → smoothed pos, velocity, covariance.",False,BLACK,False)]}),
        ("NLOS bias is systematic ⇒ modest headline gain (~4–5%) but big usability win: smoother tracks, velocity, uncertainty.", 0),
        ("", 0, {"segs":[("Twin: ",True,BLACK,False),("live tags, fading trails + true-path overlay, uncertainty rings, zone labels, live accuracy HUD, green/red LOS/NLOS anchor lines, interactive Edit-Map with hot-reload.",False,BLACK,False)]}),
    ], size=15)
    if os.path.exists(SHOT):
        s.shapes.add_picture(SHOT, Inches(7.4), Inches(1.7), width=Inches(5.3))
        tb, tf = _txbox(s, Inches(7.4), Inches(5.5), Inches(5.3), Inches(0.35))
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); _set(r, "Live 3D digital twin", size=11, italic=True, color=GREY)
builders.append(s_wd_kf)

def s_wd_calib():
    s = new("Work Done — Data-Driven NLOS Calibration (UTIL)", "Work Done")
    lw = Inches(6.5)
    tf = body(s, left=Inches(0.6), width=lw)
    bullets(tf, [
        ("", 0, {"segs":[("calibrate_from_util.py ",True,BLACK,False),("fits real range-error stats from UTIL (σ_range = σ_tdoa/√2) and patches site.json (calibrated:UTIL-IJRR2024).",False,BLACK,False)]}),
        ("", 0, {"segs":[("Finding: ",True,RED,False),("NLOS ≈ a positive bias 0.03–0.40 m, not extra noise.",False,BLACK,False)]}),
        ("Sim reads its model from config ⇒ a data change, not a code change.", 0),
    ], size=15)
    # table
    from pptx.util import Inches as In
    rows, cols = 5, 3
    tbl = s.shapes.add_table(rows, cols, Inches(7.3), Inches(1.8), Inches(5.4), Inches(2.7)).table
    data = [("Parameter","Before","After"),
            ("LOS noise σ (m)","0.08","0.054"),
            ("NLOS noise σ (m)","0.30","0.054"),
            ("NLOS bias min (m)","0.25","0.026"),
            ("NLOS bias max (m)","1.20","0.399")]
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = val
            par = cell.text_frame.paragraphs[0]
            par.alignment = PP_ALIGN.CENTER if ci else PP_ALIGN.LEFT
            run = par.runs[0]; run.font.size = Pt(13)
            run.font.bold = (ri == 0) or (ci == 2 and ri > 0)
            run.font.color.rgb = WHITE if ri == 0 else BLACK
            if ri == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xF2,0xF2,0xFA) if ri%2 else WHITE
builders.append(s_wd_calib)

def s_wd_dataset():
    s = new("Work Done — UTIL Dataset Pipeline (real UWB data)", "Work Done")
    half = (SW - Inches(1.5)) / 2
    block(s, Inches(0.6), Inches(1.5), half, Inches(1.65), "What UTIL provides",
          [("Decawave DWM1000; 4 anchor constellations; ~150 min of flights.", 0),
           ("Raw TDOA, SNR / power-difference (LOS/NLOS labels), IMU, altitude, mm-accurate Vicon ground truth.", 0)],
          bsize=12)
    block(s, Inches(0.6), Inches(3.4), half, Inches(1.8), "Parsing & alignment — dataset_source.py",
          [("Long-format event-log CSV; anchor geometry from the survey-results files.", 0),
           ("Pose column is not synced to the TDOA row ⇒ interpolate truth to each TDOA timestamp (1.2 m → ~0.09 m).", 0)],
          bsize=12)
    block(s, Inches(0.75) + half, Inches(1.5), half, Inches(2.2), "Three uses of the data",
          [("1.  Calibrate the noise / NLOS model (σ_range = σ_tdoa/√2) → patches site.json.", 0),
           ("2.  Validate the TDOA-EKF on real flights (0.18–0.33 m RMSE vs mm truth).", 0),
           ("3.  Train a 2-component Gaussian mixture (EM) → tdoa_model.json for the generative sim.", 0)],
          bsize=12)
    block(s, Inches(0.75) + half, Inches(3.95), half, Inches(1.25), "Honesty caveat",
          "Cross-radio (DWM1000 ≠ DW3000): the fitted model is a prior needing DW3000 recalibration, not a transplant.",
          bsize=12)
builders.append(s_wd_dataset)

def s_wd_ekf():
    s = new("Work Done — TDOA-EKF on Real Data + Generative Sim", "Work Done")
    tf = body(s)
    bullets(tf, [
        ("Walled-off TDOA-EKF (off the delivery critical path) validated on real UTIL flights:", 0, {"bold":True}),
        ("DatasetTdoaSource replays a trial; truth interpolated to each TDOA timestamp (1.2 m → ~0.09 m).", 1),
        ("TdoaEKF: 3D CV-EKF, h(p)=‖p−a_B‖−‖p−a_A‖ linearised each step; chi-square gate rejects NLOS outliers; Gauss–Newton cold start (no truth used).", 1),
        ("", 0, {"segs":[("Generative TDOA simulator: ",True,BLACK,False),("a learned 2-component Gaussian mixture (EM on UTIL residuals) drives any editable layout via TDOA + EKF (?map=wh-tdoa).",False,BLACK,False)]}),
    ], size=17)
builders.append(s_wd_ekf)

def s_wd_results():
    s = new("Work Done — Results", "Work Done")
    # compact 2-column table (left)
    rows = [("Configuration","Error"),
            ("Range solver, sim LOS/NLOS","~0.6–0.7 m raw"),
            ("+ Constant-velocity Kalman","~4–5% mean ↑"),
            ("TDOA-EKF, real UTIL const1/2","0.18–0.20 m"),
            ("TDOA-EKF, real UTIL const3","~0.33 m"),
            ("Generative TDOA sim (WH)","~0.5 m horiz.")]
    tw = Inches(6.1)
    tbl = s.shapes.add_table(len(rows), 2, Inches(0.5), Inches(1.7), tw, Inches(3.1)).table
    tbl.columns[0].width = Inches(4.2); tbl.columns[1].width = Inches(1.9)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci); cell.text = val
            par = cell.text_frame.paragraphs[0]; par.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
            run = par.runs[0]
            run.font.size = Pt(12); run.font.bold = (ri == 0) or (ci == 1 and ri == 3)
            run.font.color.rgb = WHITE if ri == 0 else (BLUE_DK if (ci == 1 and ri == 3) else BLACK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if ri == 0 else (RGBColor(0xF2,0xF2,0xFA) if ri % 2 else WHITE)
    # accuracy chart (right)
    if os.path.exists(CHART):
        s.shapes.add_picture(CHART, Inches(6.85), Inches(1.55), width=Inches(6.15))
    # takeaway
    tb, tf = _txbox(s, Inches(0.5), Inches(5.15), Inches(6.2), Inches(1.4))
    r = tf.paragraphs[0].add_run()
    _set(r, "Takeaway: bias dominates ⇒ NLOS weighting + calibration matter more than a heavier filter.",
         size=14, bold=True, color=BLUE_DK)
builders.append(s_wd_results)

def s_wd_test():
    s = new("Work Done — Testing & Status", "Work Done")
    lw = Inches(6.4)
    tf = body(s, left=Inches(0.6), width=lw)
    r = tf.paragraphs[0].add_run(); _set(r, "Verification", size=17, bold=True, color=BLUE_DK)
    bullets(tf, [
        ("21 automated tests (pytest), all passing.", 0, {"bold":True}),
        ("Solver < 1 cm on clean geometry; Kalman reduces mean error; calibration reproduces LOS σ + positive-only NLOS bias; EKF on synthetic + real trials.", 0),
        ("End-to-end demo: raw fix jitters, Kalman track smooth, HUD shows filtered < raw.", 0),
        ("Hardware-readiness: swapping Sim → Hardware source is the only change; engine untouched.", 0),
    ], size=14, first=False)
    rows = [("Phase","Status"),
            ("1. Foundation & framework","Done"),
            ("3. Range multilateration","Done"),
            ("4. Kalman filter fusion","Done"),
            ("5. Digital-twin upgrades","Done"),
            ("UTIL calibration","Done"),
            ("TDOA-EKF (walled-off)","Built"),
            ("Generative TDOA sim","Built"),
            ("6. Basic AI (zones/predict)","Pending")]
    tbl = s.shapes.add_table(len(rows), 2, Inches(7.25), Inches(1.55), Inches(5.45), Inches(4.4)).table
    tbl.columns[0].width = Inches(3.9); tbl.columns[1].width = Inches(1.55)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci); cell.text = val
            par = cell.text_frame.paragraphs[0]; run = par.runs[0]
            run.font.size = Pt(12); run.font.bold = (ri == 0)
            if ri == 0:
                run.font.color.rgb = WHITE; cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
            else:
                pend = val == "Pending"
                run.font.color.rgb = RED if pend else BLACK
                run.font.italic = pend
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xF2,0xF2,0xFA) if ri % 2 else WHITE
builders.append(s_wd_test)

# ------------------------------------------------------------------ WORK PLAN
def s_plan():
    s = new("Work Plan", "Work Plan")
    tb, tf = _txbox(s, Inches(0.6), Inches(1.4), SW - Inches(1.2), Inches(0.4))
    r = tf.paragraphs[0].add_run(); _set(r, "Remaining work and timeline:", size=15, italic=True, color=GREY)
    rows = [("Task","Timeline"),
            ("Phase 6 — Basic AI: geofence / zone anomaly alerts + short-horizon trajectory prediction (p + v·Δt ghost) on the Kalman velocity","Phase 1"),
            ("Hardware bring-up: implement HardwareAnchorSource for the real MaUWB_DW3000 feed; re-tune noise/power on DW3000","Phase 2"),
            ("In-situ validation: collect tape-measured, NLOS-labelled DS-TWR traces to confirm the calibration transfer direction","Phase 3"),
            ("Anchor geometry: add more / better-spread anchors to improve height observability and accuracy","Phase 4")]
    tbl = s.shapes.add_table(len(rows), 2, Inches(0.6), Inches(1.95), SW - Inches(1.2), Inches(3.6)).table
    tbl.columns[0].width = SW - Inches(1.2) - Inches(1.8); tbl.columns[1].width = Inches(1.8)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci); cell.text = val
            par = cell.text_frame.paragraphs[0]; run = par.runs[0]
            run.font.size = Pt(13); run.font.bold = (ri == 0) or ci == 1 and ri > 0
            if ri == 0:
                run.font.color.rgb = WHITE; cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
            else:
                run.font.color.rgb = BLUE_DK if ci == 1 else BLACK
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(0xF2,0xF2,0xFA) if ri % 2 else WHITE
builders.append(s_plan)

# ------------------------------------------------------------------ REFERENCES
def s_refs():
    s = new("Publication and References", "Publication & References")
    tf = body(s)
    r = tf.paragraphs[0].add_run(); _set(r, "Publication", size=17, bold=True, color=BLUE_DK)
    bullets(tf, [
        ("None yet. Target: a short paper on the data-driven NLOS calibration result (“NLOS weighting calibrated on real UWB error reduced RMSE by X%”).", 0),
    ], size=14, first=False)
    p = tf.add_paragraph(); p.space_before = Pt(8)
    r = p.add_run(); _set(r, "References", size=17, bold=True, color=BLUE_DK)
    bullets(tf, [
        ("K. Paszek, D. Grzechca, A. Becker, “Design of the UWB Positioning System Simulator for LOS/NLOS Environments,” Sensors, 21(14):4757, 2021.", 0),
        ("“NLOS-Exclusion Based UWB Positioning Algorithm,” Applied Sciences, 15:02689, 2025.", 0),
        ("“UWB Real-Time Location System Study,” Applied Sciences, 14:11005, 2024.", 0),
        ("W. Zhao, A. Goudar, X. Qiao, A. P. Schoellig, “UTIL: An Ultra-wideband TDOA Indoor Localization Dataset,” IJRR, 2024.", 0),
        ("“UWB Indoor Localization Error Modelling,” Alexandria Eng. J., 2018.", 0),
        ("Makerfabs, “MaUWB_DW3000 (ESP32 + DW3000, DS-TWR firmware),” product docs.", 0),
    ], size=13, first=False)
builders.append(s_refs)

def s_thanks():
    s = prs.slides.add_slide(BLANK)
    _slide_no["n"] += 1
    _rect(s, 0, 0, SW, SH, WHITE)
    tb, tf = _txbox(s, Inches(1), Inches(2.8), SW - Inches(2), Inches(2))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); _set(r, "Thank You", size=54, bold=True, color=BLUE)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r = p2.add_run(); _set(r, "Questions?", size=24, color=GREY)
builders.append(s_thanks)

# ------------------------------------------------------------------ BUILD
TOTAL["n"] = len(builders)
_slide_no["n"] = 0
for fn in builders:
    fn()

out = os.path.join(HERE, "UWB_RTLS_Presentation.pptx")
prs.save(out)
print("Wrote", out, "with", len(prs.slides._sldIdLst), "slides")
